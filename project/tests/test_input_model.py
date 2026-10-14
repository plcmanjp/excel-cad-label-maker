"""Synthetic Excel provenance and compatibility comparisons, not customer files."""
from dataclasses import replace
import datetime as dt
import math
from io import StringIO
from pathlib import Path
import sys
import tempfile
import unittest
from xml.etree import ElementTree as ET
import zipfile

import ezdxf
from openpyxl import Workbook
import xlwt
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from converter import LabelSettings, compute_layout, drawing_from_layout, extract_blocks_auto, save_drawing
from input_model import InputOptions, analyze, display_text, parse_range, read_workbook
from jobs import Cancelled, Cancellation


class InputModelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def workbook(self, rows, suffix=".xlsx"):
        path = self.directory / ("synthetic" + suffix)
        if suffix == ".xlsx":
            book = Workbook()
            sheet = book.active
            for row in rows:
                sheet.append(row)
            book.save(path)
            book.close()
        else:
            book = xlwt.Workbook()
            sheet = book.add_sheet("first")
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    if value is not None:
                        sheet.write(r, c, value)
            book.save(str(path))
        return path

    def test_legacy_both_formats_match_existing_reader_and_provenance(self):
        cases = [
            [["한글", "001"], [1.0, "English"], [None, None], ["NEXT", "끝"]],
            [[False, 0, True, 1], [0, False, 1, True], [None] * 4, [1.0, True, 0, False]],
            [["A\r\nB", "C\nD"], ["  SPACE  ", " "], ["tail", "tail"]],
            [["A", None], [None, "B"], [None, "C"], ["D", None]],
        ]
        for suffix in (".xlsx", ".xls"):
            for rows in cases:
                with self.subTest(suffix=suffix, rows=rows):
                    path = self.workbook(rows, suffix)
                    result = analyze(read_workbook(path))
                    self.assertEqual([list(b) for b in result.blocks], extract_blocks_auto(path))
                    self.assertTrue(all(label.addresses for label in result.labels))
                    self.assertEqual(result.cell_count, sum(map(len, result.blocks)))

    def test_sheet_range_header_and_mixed_type_scope(self):
        path = self.directory / "multiple.xlsx"
        book = Workbook()
        book.active.title = "first"
        for row in [[False, "title"], [0, "A"], [1, "B"], [None, None], [True, "C"]]:
            book.active.append(row)
        book.create_sheet("second").append(["other"])
        book.save(path)
        book.close()
        source = read_workbook(path)
        self.assertEqual(source.sheets, ("first", "second"))
        selected = analyze(source, InputOptions("A2:B3"))
        self.assertEqual(selected.blocks, (("0", "1"), ("A", "B")))
        self.assertEqual(selected.labels[0].addresses, ("A2", "A3"))
        self.assertEqual(analyze(source, InputOptions("A1:B3", True)).blocks, selected.blocks)
        self.assertEqual(analyze(read_workbook(path, "second")).blocks, (("other",),))
        with self.assertRaises(ValueError):
            read_workbook(path, "missing")

    def test_formats_diagnostics_and_source_unchanged(self):
        path = self.directory / "formats.xlsx"
        book = Workbook()
        sheet = book.active
        for address, value, fmt in [("A1", 1, "000"), ("A2", "001", "General"), ("A3", 1.234, "0.00"),
                                    ("A4", dt.datetime(2026, 10, 14, 8, 9, 10), "yyyy-mm-dd hh:mm:ss"), ("A5", .12, "0%")]:
            sheet[address] = value
            sheet[address].number_format = fmt
        sheet["B1"] = "#DIV/0!"
        sheet["B2"] = " "
        sheet["B3"] = "A\nB"
        sheet.row_dimensions[3].hidden = True
        sheet.column_dimensions["B"].hidden = True
        sheet.merge_cells("C1:D1")
        sheet["C1"] = "merged"
        sheet.auto_filter.ref = "A1:D5"
        book.save(path)
        book.close()
        initial = path.read_bytes()
        source = read_workbook(path)
        legacy = analyze(source)
        improved = analyze(source, InputOptions(mode="display"))
        self.assertEqual(improved.blocks[0], ("001", "001", "1.23", "2026-10-14 08:09:10", "0.12"))
        self.assertNotEqual(legacy.blocks[0], improved.blocks[0])
        codes = {d.code for d in improved.diagnostics}
        self.assertTrue({"error", "whitespace", "linebreak", "format-fallback", "hidden-included", "merged", "filter-included"} <= codes)
        self.assertEqual(improved.errors[0].address, "B1")
        self.assertEqual(path.read_bytes(), initial)
        self.assertTrue(source.unchanged())
        path.write_bytes(initial + b"changed")
        self.assertFalse(source.unchanged())

    def test_formula_cached_missing_error_and_incorrect_dimensions(self):
        path = self.workbook([["=2+2", "=3+3", "=1/0", "text"], ["tail"]])
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        key = "xl/worksheets/sheet1.xml"
        root = ET.fromstring(entries[key])
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        root.find("x:sheetData/x:row/x:c[@r='A1']/x:v", ns).text = "4"
        error = root.find("x:sheetData/x:row/x:c[@r='C1']", ns)
        error.set("t", "e")
        error.find("x:v", ns).text = "#DIV/0!"
        root.find("x:dimension", ns).set("ref", "A1:A1")
        entries[key] = ET.tostring(root)
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        result = analyze(read_workbook(path))
        self.assertEqual(result.blocks, (("4", "tail"), ("text",)))
        self.assertEqual({(d.code, d.address) for d in result.errors}, {("formula-missing", "B1"), ("error", "C1")})
        self.assertIn("cached-formula", {d.code for d in result.diagnostics})
        path.rename(path.with_name("unlocked.xlsx"))

    def test_xls_format_merged_hidden_and_formula_limit(self):
        path = self.directory / "format.xls"
        book = xlwt.Workbook()
        sheet = book.add_sheet("first")
        sheet.write(0, 0, 1, xlwt.easyxf(num_format_str="000"))
        sheet.write(1, 0, dt.datetime(2026, 10, 14), xlwt.easyxf(num_format_str="yyyy-mm-dd"))
        sheet.write(2, 0, xlwt.Formula("1+1"))
        sheet.write_merge(0, 0, 1, 2, "merged")
        sheet.row(1).hidden = True
        sheet.col(1).hidden = True
        book.save(str(path))
        source = read_workbook(path)
        result = analyze(source, InputOptions(mode="display"))
        self.assertEqual(result.blocks[0], ("001", "2026-10-14"))
        self.assertTrue({"xls-formula-limit", "merged", "hidden-included"} <= {d.code for d in result.diagnostics})
        self.assertFalse(any(cell.formula for cell in source.cells))

    def test_supported_format_bounds_and_negative_values(self):
        for value, fmt, expected in [(1, "000", "001"), (-1, "000", "-001"), (1.25, "0.00", "1.25"), ("001", "000", "001"),
                                     (dt.time(8, 9, 10), "hh:mm:ss", "08:09:10")]:
            self.assertEqual(display_text(value, fmt), (expected, True))
        for value, fmt in [(1.5, "000"), (10, "#,##0"), (dt.time(1, 2), "yyyy-mm-dd"), (1, "[$-ko]000")]:
            self.assertFalse(display_text(value, fmt)[1])

    def test_invalid_ranges_and_cancel_before_read(self):
        for value in ("B2:A1", "A0", "A:C", "A1:XFE2", "A1:B1048577", "A1:Z99999", "=A1", "A1 junk"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_range(value)
        token = Cancellation()
        token.cancel()
        with self.assertRaises(Cancelled):
            read_workbook(self.workbook([["A"]]), token=token)

    def test_common_layout_roundtrip_properties_and_warning_thresholds(self):
        settings = LabelSettings(cols_per_row=2, cell_w=10, min_text_height=3)
        layout = compute_layout((("가나다라마바사" * 10, "English" * 10, "001"), ("B",), ("C",)), settings)
        self.assertEqual(layout.label_bounds[2][3], -65)
        self.assertIn("small-text", {w[1] for w in layout.warnings})
        doc = drawing_from_layout(layout)
        output = self.directory / "same.dxf"
        save_drawing(doc, output)
        loaded = ezdxf.readfile(output)
        self.assertEqual(loaded.dxfversion, "AC1024")
        self.assertEqual(loaded.units, ezdxf.units.MM)
        self.assertFalse(loaded.audit().has_errors)
        texts = list(loaded.modelspace().query("TEXT"))
        boxes = list(loaded.modelspace().query("LWPOLYLINE"))
        for cell, text, box in zip(layout.cells, texts, boxes):
            self.assertEqual(text.dxf.text, cell.text)
            self.assertEqual(tuple(text.dxf.align_point)[:2], cell.center)
            self.assertAlmostEqual(text.dxf.height, cell.text_height)
            self.assertAlmostEqual(text.dxf.width, cell.width_factor)
            self.assertEqual(list(box.get_points("xy")), list(cell.corners))
            self.assertEqual((text.dxf.halign, text.dxf.valign), (1, 2))
            self.assertEqual((text.dxf.layer, box.dxf.layer), ("TEXT_DATA", "LABEL_BOX"))
        no_fit = compute_layout((("X" * 100,),), LabelSettings(auto_fit=False))
        self.assertEqual(no_fit.cells[0].text_height, 6.75)
        self.assertEqual(no_fit.warnings[0][1], "overflow")
        boundary = compute_layout((("A",),), LabelSettings(min_text_height=6.75))
        self.assertEqual(boundary.warnings, ())

    def test_nonfinite_snapshot_is_blocked(self):
        source = read_workbook(self.workbook([["A"]]))
        cell = replace(source.cells[0], value=float("inf"))
        result = analyze(replace(source, cells=(cell,)))
        self.assertEqual(result.errors[0].code, "nonfinite")

    def test_empty_formula_cache_is_distinct_from_absent_cache(self):
        path = self.workbook([["=IF(1,\"\",\"X\")", "=2+2"], ["tail", "tail"]])
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        key = "xl/worksheets/sheet1.xml"
        root = ET.fromstring(entries[key])
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        root.find("x:sheetData/x:row/x:c[@r='A1']", ns).set("t", "str")
        entries[key] = ET.tostring(root)
        with zipfile.ZipFile(path, "w") as archive:
            for name, value in entries.items():
                archive.writestr(name, value)
        result = analyze(read_workbook(path))
        self.assertEqual([(d.code, d.address) for d in result.errors], [("formula-missing", "B1")])
        self.assertIn(("formula-empty", "A1"), {(d.code, d.address) for d in result.diagnostics})

    def test_hour_patterns_preserve_single_and_double_digit_hours(self):
        for hour in (0, 8, 18):
            value = dt.time(hour, 9, 10)
            self.assertEqual(display_text(value, "h:mm"), (f"{hour}:09", True))
            self.assertEqual(display_text(value, "hh:mm"), (f"{hour:02}:09", True))
            self.assertEqual(display_text(value, "h:mm:ss"), (f"{hour}:09:10", True))

    def test_derived_overflow_and_underflow_are_refused(self):
        for kwargs in ({"cell_w": 1e308, "gap_x": 1e308}, {"cell_h": 1e-308, "text_scale": 1e-308}, {"cell_h": 1e308, "text_scale": 10}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                LabelSettings(**kwargs)
        with self.assertRaises(ValueError):
            compute_layout((("A",), ("B",), ("C",)), LabelSettings(cell_w=1e308, gap_x=0))
        with self.assertRaises(ValueError):
            compute_layout((("A", "B", "C"),), LabelSettings(cell_h=1e308, text_scale=.1))
        layout = compute_layout((tuple("A" for _ in range(10)),), LabelSettings(cell_h=1e307, text_scale=.1))
        self.assertTrue(all(math.isfinite(v) for cell in layout.cells for v in cell.center))


if __name__ == "__main__":
    unittest.main()
