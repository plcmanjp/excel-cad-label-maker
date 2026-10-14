from io import StringIO
import datetime as dt
from pathlib import Path
import sys
import tempfile
import tkinter as tk
import unittest
from unittest.mock import patch

import ezdxf
from openpyxl import Workbook
import xlwt
import zipfile
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from converter import (LabelSettings, create_drawing, extract_blocks_auto,
                       generate_dxf_string, save_drawing, select_dropped_file)


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def workbook(self, rows, name="한글 공백 샘플.xlsx"):
        path = self.directory / name
        book = Workbook()
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def test_column_order_blank_separators_and_numbers(self):
        path = self.workbook([["전원", "모터"], ["입력", 123], [None, None],
                              ["출력", "NA"], ["  ", "영문 ABC"], ["종료", None]])
        self.assertEqual(extract_blocks_auto(path),
                         [["전원", "입력"], ["출력"], ["종료"], ["모터", "123"], ["NA", "영문 ABC"]])

    def test_empty_excel_produces_no_drawing(self):
        self.assertEqual(generate_dxf_string(self.workbook([])), ("", 0))

    def test_string_api_defaults_match_ui_defaults(self):
        output, count = generate_dxf_string(self.workbook([["001"]]))
        loaded = ezdxf.read(StringIO(output))
        self.assertEqual(count, 1)
        self.assertEqual(loaded.layers.get("LABEL_BOX").dxf.color, 7)
        self.assertEqual(loaded.layers.get("TEXT_DATA").dxf.color, 7)
        self.assertEqual(loaded.styles.get("STANDARD").dxf.font, "txt.shx")

    def test_legacy_xls(self):
        path = self.directory / "구형 Excel.xls"
        book = xlwt.Workbook()
        sheet = book.add_sheet("Sheet1")
        sheet.write(0, 0, "가나다")
        sheet.write(1, 0, 2.5)
        sheet.write(3, 0, "두 번째")
        book.save(str(path))
        self.assertEqual(extract_blocks_auto(path), [["가나다", "2.5"], ["두 번째"]])

    def test_boolean_numeric_column_identity_matches_previous_reader(self):
        rows = [[False, 0, True, 1], [0, False, 1, True], ["tail"] * 4]
        expected = [["False", "False", "tail"], ["0", "0", "tail"],
                    ["True", "True", "tail"], ["1", "1", "tail"]]
        self.assertEqual(extract_blocks_auto(self.workbook(rows)), expected)
        path = self.directory / "bool-numbers.xls"
        book = xlwt.Workbook(); sheet = book.add_sheet("first")
        for r, row in enumerate(rows):
            for c, value in enumerate(row): sheet.write(r, c, value)
        book.save(str(path))
        self.assertEqual(extract_blocks_auto(path), expected)

    def test_dates_times_error_cells_and_first_sheet(self):
        path = self.workbook([[dt.datetime(2026, 1, 2, 3, 4, 5)], [dt.time(3, 4, 5)],
                              ["#DIV/0!"], ["tail"]])
        self.assertEqual(extract_blocks_auto(path),
                         [["2026-01-02 03:04:05", "03:04:05"], ["tail"]])
        path = self.directory / "date-time.xls"
        book = xlwt.Workbook(); sheet = book.add_sheet("first")
        sheet.write(0, 0, dt.datetime(2026, 1, 2, 3, 4, 5), xlwt.easyxf(num_format_str="yyyy-mm-dd hh:mm:ss"))
        sheet.write(1, 0, dt.time(3, 4, 5), xlwt.easyxf(num_format_str="hh:mm:ss"))
        book.add_sheet("ignored").write(0, 0, "ignore")
        book.save(str(path))
        self.assertEqual(extract_blocks_auto(path), [["2026-01-02 03:04:05", "03:04:05"]])

    def test_cached_formula_incorrect_dimensions_and_close(self):
        path = self.workbook([["=2+2", "first"], [None, "second"]])
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        xml = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        xml.find("x:sheetData/x:row/x:c[@r='A1']/x:v", ns).text = "4"
        xml.find("x:dimension", ns).set("ref", "A1:A1")
        entries["xl/worksheets/sheet1.xml"] = ET.tostring(xml)
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in entries.items(): archive.writestr(name, data)
        self.assertEqual(extract_blocks_auto(path), [["4"], ["first", "second"]])
        path.unlink()  # The streaming workbook must release its file on Windows.

    def test_xls_date_rounding_at_epoch_matches_previous_reader(self):
        for epoch1904 in (False, True):
            book = xlwt.Workbook(); book.dates_1904 = epoch1904
            sheet = book.add_sheet('first')
            style = xlwt.easyxf(num_format_str='yyyy-mm-dd hh:mm:ss')
            sheet.write(0, 0, 0.99999999999, style)
            sheet.write(1, 0, -1e-11, style)
            path = self.directory / f'epoch-{epoch1904}.xls'
            book.save(str(path))
            date = '1904-01-02' if epoch1904 else '1900-01-01'
            self.assertEqual(extract_blocks_auto(path), [[date + ' 00:00:00', '00:00:00']])

    def test_roundtrip_unicode_dxf_layers_entities_and_units(self):
        texts = ["한글 라벨", "한글과 English 123 😀", "긴 영문 ABC" * 10]
        doc = create_drawing([texts], LabelSettings())
        path = self.directory / "결과.dxf"
        save_drawing(doc, path)
        self.assertIn("😀", path.read_text(encoding="utf-8"))
        loaded = ezdxf.readfile(path)
        self.assertEqual(loaded.dxfversion, "AC1024")
        self.assertEqual(loaded.units, ezdxf.units.MM)
        self.assertFalse(loaded.audit().has_errors)
        self.assertEqual([e.dxf.text for e in loaded.modelspace().query("TEXT")], texts)
        self.assertEqual(len(loaded.modelspace().query("LWPOLYLINE")), 3)
        self.assertEqual(loaded.layers.get("LABEL_BOX").dxf.color, 7)
        self.assertEqual(loaded.layers.get("TEXT_DATA").dxf.color, 7)
        self.assertEqual(loaded.styles.get("STANDARD").dxf.font, "txt.shx")
        self.assertTrue(all(e.closed for e in loaded.modelspace().query("LWPOLYLINE")))

    def test_variable_height_rows_do_not_overlap(self):
        doc = create_drawing([["A", "B", "C"], ["D"], ["E"]], LabelSettings(cols_per_row=2))
        boxes = list(doc.modelspace().query("LWPOLYLINE"))
        self.assertEqual(list(boxes[3].get_points("xy"))[0], (80, 0))
        self.assertEqual(list(boxes[4].get_points("xy"))[0], (0, -65))

    def test_long_text_shrinks_and_center_alignment_is_set(self):
        settings = LabelSettings(cell_w=10)
        doc = create_drawing([["가나다라마바사" * 10, "ABCDEFGHIJ" * 10]], settings)
        for text in doc.modelspace().query("TEXT"):
            self.assertLess(text.dxf.height, 15 * .45)
            self.assertEqual(text.dxf.width, .65)
            self.assertEqual((text.dxf.halign, text.dxf.valign), (1, 2))
        doc = create_drawing([["X" * 100]], LabelSettings(auto_fit=False))
        text = next(iter(doc.modelspace().query("TEXT")))
        self.assertEqual((text.dxf.height, text.dxf.width), (6.75, 1))

    def test_invalid_settings(self):
        for key, value in [("cell_w", 0), ("cell_h", -1), ("gap_x", -1), ("gap_y", float("nan")),
                           ("text_scale", float("inf")), ("cols_per_row", 0),
                           ("cols_per_row", 2.5), ("cols_per_row", True), ("box_color", 9)]:
            with self.subTest(key=key, value=value), self.assertRaises(ValueError):
                LabelSettings(**{key: value})

    def test_corrupt_workbook_fails(self):
        path = self.directory / "壊れた.xlsx"
        path.write_text("invalid workbook", encoding="utf-8")
        with self.assertRaises(Exception):
            extract_blocks_auto(path)

    def test_dxf_tag_injection_from_multiline_cell_is_normalized(self):
        path = self.workbook([["Text\n0\nENDSEC"], [123.0]])
        output, count = generate_dxf_string(path)
        loaded = ezdxf.read(StringIO(output))
        self.assertEqual(count, 1)
        self.assertEqual([e.dxf.text for e in loaded.modelspace().query("TEXT")], ["Text 0 ENDSEC", "123"])

    def test_failed_write_preserves_existing_output(self):
        path = self.directory / "existing.dxf"
        path.write_text("previous output", encoding="utf-8")
        doc = create_drawing([["X"]], LabelSettings())
        with patch.object(doc, "write", side_effect=OSError("write failure")), self.assertRaises(OSError):
            save_drawing(doc, path)
        self.assertEqual(path.read_text(), "previous output")
        self.assertEqual(list(self.directory.glob("*.tmp")), [])
        with self.assertRaises(ValueError):
            save_drawing(doc, self.directory / "customer.xlsx")

    def test_tcl_drop_paths_spaces_braces_and_korean(self):
        interp = tk.Tcl()
        for name in ["한글 공백.xlsx", "괄호 {A} 자료.xlsx", "UPPER.XLSX"]:
            path = self.workbook([["X"]], name)
            data = interp.call("list", str(path))
            self.assertEqual(select_dropped_file(interp, data), path.resolve())

    def test_reject_multiple_wrong_missing_and_excel_temp_files(self):
        interp = tk.Tcl()
        path = self.workbook([["X"]])
        for data in ["", interp.call("list", str(path), str(path)),
                     interp.call("list", str(self.directory)),
                     interp.call("list", str(self.directory / "missing.xlsx")),
                     interp.call("list", str(self.workbook([["X"]], "~$temp.xlsx"))),
                     interp.call("list", str(self.directory / "drawing.dxf"))]:
            with self.subTest(data=data), self.assertRaises(ValueError):
                select_dropped_file(interp, data)


if __name__ == "__main__":
    unittest.main()
