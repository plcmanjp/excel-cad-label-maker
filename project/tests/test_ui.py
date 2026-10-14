"""Real Tk widgets, synthetic drop data; does not claim Explorer mouse QC."""
import ast
from pathlib import Path
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from types import SimpleNamespace
from threading import Event

from openpyxl import Workbook
from tkinterdnd2 import COPY, REFUSE_DROP, TkinterDnD
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from excel_cad_label import App


class UiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.file = Path(self.temp.name) / "드래그 공백 파일.xlsx"
        book = Workbook()
        book.active.append(["TEST"])
        book.save(self.file)
        book.close()
        self.root = TkinterDnD.Tk()
        self.root.withdraw()
        self.addCleanup(self.root.destroy)
        self.app = App(self.root, check_updates=False, settings_dir=self.file.parent / "settings")

    def pump(self, predicate, timeout=5):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            self.root.update()
            if predicate():
                return
            time.sleep(.005)
        self.fail("Tk job did not finish: " + self.app.status.get())

    def test_drop_selects_file_without_browse_and_enables_generate(self):
        with patch("excel_cad_label.filedialog.askopenfilename", side_effect=AssertionError("browse called")):
            event = SimpleNamespace(data=self.root.tk.call("list", str(self.file)))
            self.assertEqual(self.app.on_drop(event), COPY)
        self.assertEqual(self.app.excel_path.get(), str(self.file.resolve()))
        self.pump(lambda: not self.app.busy)
        self.assertEqual(str(self.app.run_button.cget("state")), "normal")
        self.assertIn(self.file.name, self.app.drop_area.cget("text"))

    def test_bad_drop_preserves_selection_and_clear_disables_generate(self):
        self.app.select_file(self.file)
        self.assertEqual(self.app.on_drop(SimpleNamespace(data="bad.txt")), REFUSE_DROP)
        self.assertEqual(self.app.excel_path.get(), str(self.file.resolve()))
        self.app.clear_file()
        self.assertEqual(self.app.excel_path.get(), "")
        self.assertEqual(str(self.app.run_button.cget("state")), "disabled")

    def test_browse_and_drop_share_selection(self):
        with patch("excel_cad_label.filedialog.askopenfilename", return_value=str(self.file)):
            self.app.on_browse()
        self.assertEqual(self.app.excel_path.get(), str(self.file.resolve()))

    def test_generate_saves_dxf_after_drop_without_open_dialog(self):
        import ezdxf
        output = Path(self.temp.name) / "output.dxf"
        self.app.on_drop(SimpleNamespace(data=self.root.tk.call("list", str(self.file))))
        self.pump(lambda: not self.app.busy)
        with patch("excel_cad_label.filedialog.askopenfilename", side_effect=AssertionError("browse called")), \
             patch("excel_cad_label.filedialog.asksaveasfilename", return_value=str(output)), \
             patch("excel_cad_label.messagebox.showinfo"):
            self.app.on_generate()
            self.pump(lambda: not self.app.busy)
        self.assertEqual(next(iter(ezdxf.readfile(output).modelspace().query("TEXT"))).dxf.text, "TEST")

    def test_registered_drops_on_entire_window(self):
        def check(widget):
            self.assertTrue(widget.dnd_bind("<<Drop>>"))
            for child in widget.winfo_children():
                check(child)
        check(self.root)

    def test_corrupt_drop_restores_previous_input_and_result(self):
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        previous = self.app.analysis
        bad = self.file.parent / "corrupt.xlsx"
        bad.write_bytes(b"not Excel")
        self.app.select_file(bad)
        self.pump(lambda: not self.app.busy)
        self.assertIs(self.app.analysis, previous)
        self.assertEqual(self.app.excel_path.get(), str(self.file.resolve()))
        self.assertIn("유지", self.app.status.get())

    def test_warnings_confirm_once_and_change_invalidates_confirmation(self):
        book = Workbook()
        book.active.append(["A\nB"])
        book.save(self.file)
        book.close()
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        self.assertEqual(str(self.app.run_button.cget("state")), "disabled")
        self.app.confirm_warnings()
        self.assertEqual(str(self.app.run_button.cget("state")), "normal")
        self.app.vars["cell_w"].set("70")
        self.pump(lambda: not self.app.busy and self.app.debounce_id is None)
        self.assertFalse(self.app.confirmed)
        self.assertEqual(str(self.app.run_button.cget("state")), "disabled")

    def test_errors_block_save_and_keep_source_address(self):
        book = Workbook()
        book.active.append(["#DIV/0!"])
        book.active.append(["tail"])
        book.save(self.file)
        book.close()
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        self.assertIn("저장 차단", self.app.save_reason.get())
        self.assertEqual(self.app.analysis.errors[0].address, "A1")
        with patch("excel_cad_label.filedialog.asksaveasfilename", side_effect=AssertionError("blocked save")):
            self.app.on_generate()

    def test_style_change_reuses_snapshot_and_page_selection(self):
        book = Workbook()
        for index in range(220):
            book.active.append([f"L{index}"])
            book.active.append([None])
        book.save(self.file)
        book.close()
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        source = self.app.snapshot
        with patch("label_ui.read_workbook", side_effect=AssertionError("Excel reread")):
            self.app.vars["cell_w"].set("72")
            self.pump(lambda: not self.app.busy and self.app.debounce_id is None)
        self.assertIs(self.app.snapshot, source)
        self.assertEqual(len(self.app.data_tree.get_children()), 100)
        self.app.select_label(205)
        self.assertEqual(self.app.page, 2)
        self.assertEqual(self.app.preview.selected, 205)
        self.assertIn("라벨 206", self.app.selection_info.get())

    def test_source_change_blocks_existing_output_and_reload_recovers(self):
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        book = Workbook()
        book.active.append(["NEW"])
        book.save(self.file)
        book.close()
        output = self.file.parent / "existing.dxf"
        output.write_bytes(b"existing")
        with patch("excel_cad_label.filedialog.asksaveasfilename", return_value=str(output)):
            self.app.on_generate()
            self.pump(lambda: not self.app.busy)
        self.assertEqual(output.read_bytes(), b"existing")
        self.assertTrue(self.app.source_changed)
        self.app.reload()
        self.pump(lambda: not self.app.busy)
        self.assertEqual(self.app.analysis.blocks, (("NEW",),))
        self.assertFalse(self.app.source_changed)

    def test_cancel_save_dialog_and_duplicate_save_are_safe(self):
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        with patch("excel_cad_label.filedialog.asksaveasfilename", return_value=""):
            self.app.on_generate()
        self.assertFalse(self.app.busy)
        output = self.file.parent / "output.dxf"
        started, release = Event(), Event()
        import label_ui
        actual = label_ui.drawing_from_layout
        def delayed(*args):
            started.set()
            if not release.wait(2):
                raise AssertionError("save timeout")
            return actual(*args)
        with patch("excel_cad_label.filedialog.asksaveasfilename", return_value=str(output)) as dialog, patch("label_ui.drawing_from_layout", side_effect=delayed):
            self.app.on_generate()
            self.assertTrue(started.wait(2))
            self.app.on_generate()
            self.assertEqual(dialog.call_count, 1)
            self.app.cancel()
            release.set()
            self.pump(lambda: not self.app.busy)
        self.assertFalse(output.exists())

    def test_old_file_result_and_setting_change_during_read(self):
        second = self.file.parent / "B.xlsx"
        book = Workbook()
        book.active.append(["SECOND"])
        book.save(second)
        book.close()
        started, release = Event(), Event()
        import label_ui
        actual = label_ui.read_workbook
        def delayed(path, sheet=None, token=None):
            # Windows TEMP may use an 8.3 alias; selection resolves that alias.
            if Path(path).resolve() == self.file.resolve():
                started.set()
                if not release.wait(2):
                    raise AssertionError("read timeout")
            return actual(path, sheet, token)
        with patch("label_ui.read_workbook", side_effect=delayed):
            self.app.select_file(self.file)
            self.assertTrue(started.wait(2))
            self.app.select_file(second)
            self.app.vars["cell_w"].set("71")
            release.set()
            self.pump(lambda: not self.app.busy and self.app.debounce_id is None)
        self.assertEqual(self.app.snapshot.path, second.resolve())
        self.assertEqual(self.app.analysis.blocks, (("SECOND",),))
        self.assertEqual(self.app.layout.settings.cell_w, 71)

    def test_invalid_setting_field_focus_and_restore_last_settings(self):
        for value in ("0", "-1", "NaN", "Infinity", "", "abc"):
            self.app.vars["cell_w"].set(value)
            self.assertIsNone(self.app.parse_settings())
            self.assertEqual(self.app.invalid_field, "cell_w")
            self.assertIn("가로폭", self.app.form_error.get())
        self.app.vars["cell_w"].set("74")
        self.app.remember()
        from settings_store import SettingsStore
        self.assertEqual(SettingsStore(self.file.parent / "settings").last["cell_w"], 74)

    def test_preset_ui_dirty_restore_rename_delete(self):
        self.app.vars["cell_w"].set("80")
        self.assertEqual(self.app.preset_state.get(), "수정됨")
        with patch("label_ui.simpledialog.askstring", return_value="규격 A"):
            self.app.save_preset()
        self.assertEqual(self.app.preset_state.get(), "선택 설정과 같음")
        self.app.vars["cell_w"].set("85")
        self.app.load_preset()
        self.assertEqual(self.app.vars["cell_w"].get(), "80")
        with patch("label_ui.simpledialog.askstring", return_value="규격 B"):
            self.app.rename_preset()
        self.assertIn("규격 B", self.app.store.presets)
        with patch("excel_cad_label.messagebox.askyesno", return_value=False):
            self.app.delete_preset()
        self.assertIn("규격 B", self.app.store.presets)
        with patch("excel_cad_label.messagebox.askyesno", return_value=True):
            self.app.delete_preset()
        self.assertEqual(self.app.store.presets, {})

    def test_open_missing_cad_keeps_success_and_example_is_20_labels_60_cells(self):
        example = self.file.parent / "example.xlsx"
        with patch("excel_cad_label.filedialog.asksaveasfilename", return_value=str(example)):
            self.app.make_example()
        self.pump(lambda: not self.app.busy)
        self.assertEqual((len(self.app.analysis.labels), self.app.analysis.cell_count), (20, 60))
        self.app.output = self.file.parent / "saved.dxf"
        with patch("label_ui.os.startfile", side_effect=OSError()):
            self.app.open_output()
        self.assertIn("저장은 완료", self.app.status.get())

    def test_shutdown_cancels_workers_and_after_callbacks(self):
        started, release = Event(), Event()
        import label_ui
        actual = label_ui.read_workbook
        def delayed(*args):
            started.set()
            release.wait(2)
            return actual(*args)
        with patch("label_ui.read_workbook", side_effect=delayed):
            self.app.select_file(self.file)
            self.assertTrue(started.wait(2))
            self.app.shutdown()
            release.set()
        self.assertTrue(self.app.closed)
        self.assertTrue(self.app.jobs.closed)
        self.assertIsNone(self.app.after_id)
        self.assertIsNone(self.app.debounce_id)

    def test_all_save_exit_paths_apply_edits_made_during_save(self):
        for failing in (False, True):
            with self.subTest(failing=failing):
                self.app.select_file(self.file)
                self.pump(lambda: not self.app.busy)
                started, release = Event(), Event()
                import label_ui
                actual = label_ui.drawing_from_layout
                def delayed(*args):
                    started.set()
                    release.wait(2)
                    if failing:
                        raise OSError("synthetic write failure")
                    return actual(*args)
                output = self.file.parent / "preserved.dxf"
                output.write_bytes(b"preserve")
                with patch("excel_cad_label.filedialog.asksaveasfilename", return_value=str(output)), patch("label_ui.drawing_from_layout", side_effect=delayed):
                    self.app.on_generate()
                    self.assertTrue(started.wait(2))
                    self.app.vars["cell_w"].set("75" if failing else "76")
                    if not failing:
                        self.app.cancel()
                    release.set()
                    self.pump(lambda: not self.app.busy and self.app.debounce_id is None)
                self.assertEqual(output.read_bytes(), b"preserve")
                self.assertEqual(self.app.layout.settings.cell_w, 75 if failing else 76)
                self.assertEqual(str(self.app.run_button.cget("state")), "normal")

    def test_diagnostic_pages_reach_error_201(self):
        book = Workbook()
        for _ in range(201):
            book.active.append(["#DIV/0!"])
        book.active.append(["tail"])
        book.save(self.file)
        book.close()
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        self.assertEqual(len(self.app.diag_tree.get_children()), 200)
        self.app.change_diagnostic_page(1)
        self.assertEqual(len(self.app.diag_tree.get_children()), 1)
        self.app.diag_tree.selection_set("200")
        self.app.diagnostic_selected()
        self.assertIn("A201", self.app.selection_info.get())

    def test_single_long_label_preview_and_detail_are_bounded(self):
        book = Workbook()
        for index in range(3000):
            book.active.append([f"LINE {index}"])
        book.save(self.file)
        book.close()
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        self.app.preview.draw()
        self.assertLess(len(self.app.preview.canvas.find_all()), 1201)
        self.assertIn("전체 3,000칸", self.app.data_tree.item("0", "values")[2])
        self.app.select_label(0)
        self.assertLess(len(self.app.selection_info.get()), 400)
        self.app.show_label_details()
        dialogs = [child for child in self.root.winfo_children() if child.winfo_class() == "Toplevel"]
        self.assertEqual(len(dialogs), 1)
        trees = [child for child in dialogs[0].winfo_children() if child.winfo_class() == "Treeview"]
        self.assertEqual(len(trees[0].get_children()), 100)
        dialogs[0].destroy()

    def test_example_cannot_overwrite_pending_first_or_replacement_input(self):
        original = self.file.read_bytes()
        started, release = Event(), Event()
        import label_ui
        actual = label_ui.read_workbook
        def delayed(*args):
            started.set()
            release.wait(2)
            return actual(*args)
        with patch("label_ui.read_workbook", side_effect=delayed):
            self.app.select_file(self.file)
            self.assertTrue(started.wait(2))
            with patch("excel_cad_label.filedialog.asksaveasfilename", return_value=str(self.file)):
                self.app.make_example()
            self.assertEqual(self.file.read_bytes(), original)
            release.set()
            self.pump(lambda: not self.app.busy)

        second = self.file.parent / "B.xlsx"
        second.write_bytes(original)
        started.clear()
        release.clear()
        with patch("label_ui.read_workbook", side_effect=delayed):
            self.app.select_file(second)
            self.assertTrue(started.wait(2))
            with patch("excel_cad_label.filedialog.asksaveasfilename", return_value=str(second)):
                self.app.make_example()
            self.assertEqual(second.read_bytes(), original)
            release.set()
            self.pump(lambda: not self.app.busy)

    def test_native_control_names_and_keyboard_tabs_preserve_values(self):
        import ctypes
        from ctypes import wintypes
        function = ctypes.WinDLL("user32").GetWindowTextW
        function.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
        buffer = ctypes.create_unicode_buffer(256)
        function(self.app.browse_button.winfo_id(), buffer, 256)
        self.assertEqual(buffer.value, "파일 찾기")
        self.assertEqual(self.app.parse_settings().cell_w, 60)
        self.assertEqual(self.app.switch_tab(1), "break")
        self.assertEqual(self.app.tabs.index(self.app.tabs.select()), 1)
        self.app.switch_tab(-1)
        self.assertEqual(self.app.tabs.index(self.app.tabs.select()), 0)
        self.app.preview.schedule()
        self.app.preview.draw()
        self.assertIsNone(self.app.preview.pending)

    def test_no_prohibited_ui_punctuation(self):
        forbidden = set("\u2014\u2013\u2015\u00b7\u30fb\u2022")
        for source in (Path(__file__).resolve().parents[1] / "src").glob("*.py"):
            for node in ast.walk(ast.parse(source.read_text(encoding="utf-8"))):
                if isinstance(node, ast.Constant) and isinstance(node.value, str):
                    self.assertFalse(forbidden.intersection(node.value), (source.name, node.lineno))

    def test_preview_refits_after_hidden_tab_is_shown_and_keeps_manual_zoom(self):
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        initial = self.app.preview.scale
        self.root.deiconify()
        self.app.tabs.select(self.app.preview)
        self.root.update()
        self.assertGreater(self.app.preview.scale, initial)
        self.app.preview.zoom(1.35)
        manual = self.app.preview.scale
        self.root.geometry("1000x650")
        self.root.update()
        self.assertEqual(self.app.preview.scale, manual)
        self.app.preview.fit()
        self.assertTrue(self.app.preview.fit_mode)

    def test_hidden_selected_label_fits_when_preview_is_first_shown(self):
        self.app.select_file(self.file)
        self.pump(lambda: not self.app.busy)
        self.app.select_label(0, from_list=True)
        self.assertTrue(self.app.preview.selected_fit_pending)
        self.root.deiconify()
        self.app.tabs.select(self.app.preview)
        self.root.update()
        self.assertFalse(self.app.preview.selected_fit_pending)
        width, height = self.app.preview.viewport()
        self.assertAlmostEqual(self.app.preview.scale, min((width - 60) / 60, (height - 60) / 15, 15))

    def test_dialog_reentry_and_shortcuts_are_blocked_until_cancel_or_failure(self):
        def nested_dialog(**kwargs):
            self.assertTrue(self.app.modal_open)
            self.app.on_browse()
            self.app.on_generate()
            self.app.make_example()
            self.app._shortcut(self.app.on_browse)
            self.app.save_preset()
            return ""
        with patch("excel_cad_label.filedialog.askopenfilename", side_effect=nested_dialog) as dialog:
            self.app.on_browse()
            self.assertEqual(dialog.call_count, 1)
        self.assertFalse(self.app.modal_open)
        with patch("excel_cad_label.filedialog.askopenfilename", side_effect=OSError("dialog failure")):
            with self.assertRaises(OSError):
                self.app.on_browse()
        self.assertFalse(self.app.modal_open)


if __name__ == "__main__":
    unittest.main()
