"""Immutable Excel input with provenance, explicit diagnostics and bounded formats."""
from dataclasses import dataclass
from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
from itertools import zip_longest
import math
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
import xlrd

from converter import _label_text, validate_excel_path
from jobs import Cancellation

MAX_CELLS = 250_000
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


@dataclass(frozen=True)
class Diagnostic:
    severity: str
    code: str
    message: str
    address: str = ""
    label: int | None = None


@dataclass(frozen=True)
class Cell:
    row: int
    column: int
    value: object
    kind: str
    number_format: str = "General"
    formula: bool = False

    @property
    def address(self):
        return f"{get_column_letter(self.column)}{self.row}"


@dataclass(frozen=True)
class WorkbookInput:
    path: Path
    fingerprint: str
    sheets: tuple
    sheet: str
    cells: tuple
    rows: int
    columns: int
    merged: tuple = ()
    hidden_rows: frozenset = frozenset()
    hidden_columns: frozenset = frozenset()
    filtered: bool = False
    legacy_xls: bool = False

    def unchanged(self):
        try:
            return sha256(self.path.read_bytes()).hexdigest() == self.fingerprint
        except OSError:
            return False


@dataclass(frozen=True)
class InputOptions:
    cell_range: str = ""
    skip_header: bool = False
    mode: str = "legacy"

    def __post_init__(self):
        if self.mode not in {"legacy", "display"} or type(self.skip_header) is not bool:
            raise ValueError("입력 해석 설정을 확인해 주세요.")
        if self.cell_range:
            parse_range(self.cell_range)


@dataclass(frozen=True)
class Label:
    texts: tuple
    addresses: tuple
    origin: str = "excel"


@dataclass(frozen=True)
class Analysis:
    source: WorkbookInput | None
    options: InputOptions
    applied_range: str
    labels: tuple
    diagnostics: tuple

    @property
    def blocks(self):
        return tuple(label.texts for label in self.labels)

    @property
    def cell_count(self):
        return sum(len(label.texts) for label in self.labels)

    @property
    def errors(self):
        return tuple(d for d in self.diagnostics if d.severity == "error")


def parse_range(value):
    value = value.strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*(?::[A-Z]{1,3}[1-9][0-9]*)?", value):
        raise ValueError("셀 범위는 A1 또는 A1:C20 형식으로 입력해 주세요.")
    bounds = range_boundaries(value)
    c1, r1, c2, r2 = bounds
    if c2 < c1 or r2 < r1 or c2 > 16384 or r2 > 1048576:
        raise ValueError("셀 범위의 시작과 끝을 확인해 주세요.")
    if (c2 - c1 + 1) * (r2 - r1 + 1) > MAX_CELLS:
        raise ValueError(f"선택 범위는 {MAX_CELLS:,}셀 이내로 지정해 주세요.")
    return bounds


def _xlsx_metadata(data, path, token):
    merged, hidden_rows, hidden_columns, filtered = [], set(), set(), False
    empty_formula_results = set()
    current_cell = None
    with ZipFile(BytesIO(data)) as archive, archive.open(path) as stream:
        for event, item in ET.iterparse(stream, events=("start", "end")):
            token.check()
            if event == "start":
                if item.tag == NS + "c":
                    current_cell = {"address": item.get("r"), "type": item.get("t"), "formula": False, "value": False, "text": None}
                continue
            if current_cell is not None:
                if item.tag == NS + "f":
                    current_cell["formula"] = True
                elif item.tag == NS + "v":
                    current_cell["value"] = True
                    current_cell["text"] = item.text
                elif item.tag == NS + "c":
                    if current_cell["formula"] and current_cell["value"] and current_cell["type"] == "str" and not current_cell["text"]:
                        empty_formula_results.add(current_cell["address"])
                    current_cell = None
            if item.tag == NS + "row" and item.get("hidden") in {"1", "true"}:
                hidden_rows.add(int(item.get("r")))
            elif item.tag == NS + "col" and item.get("hidden") in {"1", "true"}:
                hidden_columns.update(range(int(item.get("min")), int(item.get("max")) + 1))
            elif item.tag == NS + "mergeCell":
                merged.append(item.get("ref"))
            elif item.tag == NS + "autoFilter" and item.get("ref"):
                filtered = True
            item.clear()
    return (tuple(merged), frozenset(hidden_rows), frozenset(hidden_columns), filtered), empty_formula_results


def read_workbook(value, sheet_name=None, token=None):
    token = token or Cancellation()
    token.check()
    path = validate_excel_path(value)
    before = path.stat()
    data = path.read_bytes()
    after = path.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise ValueError("Excel 파일이 읽는 동안 변경되었습니다. 다시 선택해 주세요.")
    fingerprint = sha256(data).hexdigest()
    cells, rows, columns = [], 0, 0
    if path.suffix.lower() == ".xlsx":
        with BytesIO(data) as cached_stream, BytesIO(data) as formula_stream:
            cached = load_workbook(cached_stream, read_only=True, data_only=True, keep_links=False)
            try:
                formulas = load_workbook(formula_stream, read_only=True, data_only=False, keep_links=False)
                try:
                    names = tuple(cached.sheetnames)
                    name = sheet_name or names[0]
                    if name not in names:
                        raise ValueError("선택한 시트가 없습니다. 시트를 다시 선택해 주세요.")
                    sheet, f_sheet = cached[name], formulas[name]
                    metadata, empty_formula_results = _xlsx_metadata(data, sheet._worksheet_path, token)
                    sheet.reset_dimensions()
                    f_sheet.reset_dimensions()
                    for r, (row, f_row) in enumerate(zip_longest(sheet.iter_rows(), f_sheet.iter_rows(), fillvalue=()), 1):
                        token.check()
                        rows = r
                        columns = max(columns, len(row), len(f_row))
                        if rows * columns > MAX_CELLS:
                            raise ValueError(f"시트는 {MAX_CELLS:,}셀 이내로 나누어 저장해 주세요.")
                        for c, (cell, f_cell) in enumerate(zip_longest(row, f_row), 1):
                            value = getattr(cell, "value", None)
                            is_formula = getattr(f_cell, "data_type", "n") == "f"
                            if is_formula and f"{get_column_letter(c)}{r}" in empty_formula_results:
                                value = ""
                            kind = "error" if getattr(cell, "data_type", "n") == "e" else "formula-missing" if is_formula and value is None else "blank" if value is None else "value"
                            if value is None and not is_formula and getattr(f_cell, "data_type", "n") in {"s", "str", "inlineStr"}:
                                value, kind = "", "empty-string"
                            cells.append(Cell(r, c, value, kind, getattr(cell, "number_format", "General") or "General", is_formula))
                finally:
                    formulas.close()
            finally:
                cached.close()
        return WorkbookInput(path, fingerprint, names, name, tuple(cells), rows, columns, *metadata)
    book = xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=True)
    try:
        names = tuple(book.sheet_names())
        name = sheet_name or names[0]
        if name not in names:
            raise ValueError("선택한 시트가 없습니다. 시트를 다시 선택해 주세요.")
        sheet = book.sheet_by_name(name)
        rows, columns = sheet.nrows, sheet.ncols
        if rows * columns > MAX_CELLS:
            raise ValueError(f"시트는 {MAX_CELLS:,}셀 이내로 나누어 저장해 주세요.")
        for r in range(rows):
            token.check()
            for c in range(columns):
                cell = sheet.cell(r, c)
                value = cell.value
                kind = "error" if cell.ctype == xlrd.XL_CELL_ERROR else "blank" if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK} else "value"
                if kind == "blank":
                    value = None
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        parsed = xlrd.xldate_as_datetime(value, book.datemode)
                        anchor = (1904, 1, 1) if book.datemode else (1899, 12, 31)
                        value = parsed.time() if (parsed.year, parsed.month, parsed.day) == anchor else parsed
                    except (ValueError, OverflowError):
                        pass
                xf = book.xf_list[sheet.cell_xf_index(r, c)]
                fmt = book.format_map[xf.format_key].format_str
                cells.append(Cell(r + 1, c + 1, value, kind, fmt))
        merged = tuple(f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2)}{r2}" for r1, r2, c1, c2 in sheet.merged_cells)
        hidden_rows = frozenset(r + 1 for r, info in sheet.rowinfo_map.items() if info.hidden)
        hidden_columns = frozenset(c + 1 for c, info in sheet.colinfo_map.items() if info.hidden)
        return WorkbookInput(path, fingerprint, names, name, tuple(cells), rows, columns, merged, hidden_rows, hidden_columns, False, True)
    finally:
        book.release_resources()


DATE_FORMATS = {
    "yyyy-mm-dd": "%Y-%m-%d", "yyyy/mm/dd": "%Y/%m/%d",
    "yyyy-mm-dd hh:mm": "%Y-%m-%d %H:%M", "yyyy-mm-dd hh:mm:ss": "%Y-%m-%d %H:%M:%S",
    "hh:mm": "%H:%M", "hh:mm:ss": "%H:%M:%S", "h:mm": "%H:%M", "h:mm:ss": "%H:%M:%S",
}


def display_text(value, number_format):
    fmt = number_format.strip()
    if fmt in {"General", "@"} or isinstance(value, (str, bool)) or value is None:
        return _label_text(value), True
    if isinstance(value, (date, datetime, time)):
        pattern = DATE_FORMATS.get(fmt.lower())
        if pattern:
            if isinstance(value, time) and "%Y" in pattern:
                return _label_text(value), False
            if isinstance(value, date) and not isinstance(value, datetime) and "%H" in pattern:
                return _label_text(value), False
            if fmt.lower() in {"h:mm", "h:mm:ss"}:
                suffix = f":{value.second:02}" if fmt.lower() == "h:mm:ss" else ""
                return f"{value.hour}:{value.minute:02}{suffix}", True
            return value.strftime(pattern), True
    elif type(value) in {int, float} and math.isfinite(value):
        if re.fullmatch(r"0{1,15}", fmt) and float(value).is_integer():
            sign = "-" if value < 0 else ""
            return sign + str(abs(int(value))).zfill(len(fmt)), True
        if re.fullmatch(r"0\.0{1,10}", fmt):
            return format(value, f".{len(fmt) - 2}f"), True
    return _label_text(value), False


def analyze(source, options=None, token=None):
    options, token = options or InputOptions(), token or Cancellation()
    c1, r1, c2, r2 = parse_range(options.cell_range) if options.cell_range.strip() else (1, 1, max(1, source.columns), max(1, source.rows))
    if options.skip_header:
        r1 += 1
    selected = [cell for cell in source.cells if r1 <= cell.row <= r2 and c1 <= cell.column <= c2]
    diagnostics = []
    if source.legacy_xls:
        diagnostics.append(Diagnostic("warning", "xls-formula-limit", ".xls는 수식 출처와 결과 누락을 구분하지 못합니다. Excel에서 계산하고 저장한 값을 확인해 주세요."))
    if any(r1 <= row <= r2 for row in source.hidden_rows) or any(c1 <= col <= c2 for col in source.hidden_columns):
        diagnostics.append(Diagnostic("warning", "hidden-included", "숨김 행과 열의 데이터도 포함합니다. 자동 제외하지 않습니다."))
    for merged in source.merged:
        mc1, mr1, mc2, mr2 = range_boundaries(merged)
        if mc1 <= c2 and mc2 >= c1 and mr1 <= r2 and mr2 >= r1:
            diagnostics.append(Diagnostic("warning", "merged", "병합 영역은 기준 셀의 값만 읽습니다. 나머지 빈 셀은 라벨 구분에 사용됩니다.", merged))
    if source.filtered:
        diagnostics.append(Diagnostic("warning", "filter-included", "필터가 있는 시트도 데이터와 숨김 행을 포함합니다. 자동 제외하지 않습니다."))
    columns = {}
    for cell in selected:
        columns.setdefault(cell.column, []).append(cell)
    labels, address_labels = [], {}
    for column in sorted(columns):
        token.check()
        active_texts, active_addresses, boolean_numbers = [], [], {}
        seen_bool, seen_number = False, False

        def flush():
            if active_texts:
                number = len(labels)
                labels.append(Label(tuple(active_texts), tuple(active_addresses)))
                address_labels.update((address, number) for address in active_addresses)
                active_texts.clear()
                active_addresses.clear()

        previous_row = r1 - 1
        for cell in columns[column]:
            token.check()
            if cell.row > previous_row + 1:
                flush()
            previous_row = cell.row
            value = cell.value
            if cell.kind in {"error", "formula-missing"}:
                message = "Excel 오류 셀입니다. 원본에서 오류를 수정하고 저장해 주세요." if cell.kind == "error" else "수식의 저장된 결과가 없습니다. Excel에서 계산하고 저장해 주세요."
                diagnostics.append(Diagnostic("error", cell.kind, message, cell.address))
                text = ""
            elif type(value) is float and not math.isfinite(value):
                diagnostics.append(Diagnostic("error", "nonfinite", "유한하지 않은 숫자는 출력할 수 없습니다.", cell.address))
                text = ""
            else:
                if cell.formula:
                    diagnostics.append(Diagnostic("warning", "cached-formula", "저장된 수식 결과를 사용합니다. 최신 계산값인지는 확인할 수 없습니다.", cell.address))
                    if value == "":
                        diagnostics.append(Diagnostic("warning", "formula-empty", "수식의 저장된 결과가 빈 문자열이어서 라벨을 구분합니다.", cell.address))
                elif cell.kind == "empty-string":
                    diagnostics.append(Diagnostic("warning", "empty-string", "빈 문자열은 빈 출력으로 처리되어 라벨을 구분합니다.", cell.address))
                if type(value) in {bool, int, float} and value in (0, 1):
                    seen_bool |= type(value) is bool
                    seen_number |= type(value) is not bool
                    if options.mode == "legacy":
                        value = boolean_numbers.setdefault(value, value)
                if options.mode == "display":
                    text, supported = display_text(value, cell.number_format)
                    if not supported:
                        diagnostics.append(Diagnostic("warning", "format-fallback", f"미지원 표시 형식 {cell.number_format!r}: 대체 출력 {text!r}를 확인해 주세요.", cell.address))
                else:
                    text = _label_text(value)
                    if cell.number_format not in {"General", "@"} and type(value) is not str:
                        formatted, supported = display_text(value, cell.number_format)
                        if not supported or formatted != text:
                            diagnostics.append(Diagnostic("warning", "legacy-format", f"기존 해석은 표시 형식을 적용하지 않습니다. 출력 {text!r}를 확인해 주세요.", cell.address))
                if isinstance(value, str):
                    if value and not value.strip():
                        diagnostics.append(Diagnostic("warning", "whitespace", "공백 문자열은 빈 출력으로 처리되어 라벨을 구분합니다.", cell.address))
                    elif value.strip() != value:
                        diagnostics.append(Diagnostic("warning", "trimmed", f"양끝 공백을 정리한 출력 {text!r}를 확인해 주세요.", cell.address))
                    if "\n" in value or "\r" in value:
                        diagnostics.append(Diagnostic("warning", "linebreak", f"셀 줄바꿈을 한 줄로 연결한 출력 {text!r}를 확인해 주세요.", cell.address))
            if text:
                active_texts.append(text)
                active_addresses.append(cell.address)
            else:
                flush()
        flush()
        if seen_bool and seen_number:
            diagnostics.append(Diagnostic("warning", "mixed-boolean", f"{get_column_letter(column)}열의 논리값과 숫자가 섞여 있습니다. 선택한 해석 모드의 출력 문자를 확인해 주세요."))
    from dataclasses import replace
    diagnostics = tuple(replace(d, label=address_labels.get(d.address)) for d in diagnostics)
    applied = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}" if r1 <= r2 else "제목 행 제외 후 빈 범위"
    return Analysis(source, options, applied, tuple(labels), diagnostics)
