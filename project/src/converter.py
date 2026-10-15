"""Recovered column-oriented label conversion. Original credit: 직장인."""
from dataclasses import dataclass
from io import StringIO
import math
import os
from pathlib import Path
import tempfile
import unicodedata

import ezdxf
from ezdxf.enums import TextEntityAlignment
from openpyxl import load_workbook
import xlrd

CAD_COLORS = {
    "기본(검정)": 7, "빨강 (Red, 1)": 1,
    "노랑 (Yellow, 2)": 2, "초록 (Green, 3)": 3,
    "하늘 (Cyan, 4)": 4, "파랑 (Blue, 5)": 5,
    "분홍 (Magenta, 6)": 6,
}
CAD_FONTS = {
    "맑은 고딕 (Malgun Gothic)": ("MALGUN", "malgun.ttf"),
    "굴림 (Gulim)": ("GULIM", "gulim.ttc"),
    "돋움 (Dotum)": ("DOTUM", "dotum.ttc"),
    "바탕 (Batang)": ("BATANG", "batang.ttc"),
    "Arial": ("ARIAL", "arial.ttf"),
    "AutoCAD 표준 (Standard)": ("STANDARD", "txt.shx"),
}
DEFAULT_FONT = "Arial"


def validate_excel_path(value):
    path = Path(value).expanduser()
    if path.suffix.lower() not in {".xlsx", ".xls"}:
        raise ValueError("Excel 파일(.xlsx 또는 .xls)을 선택해 주세요.")
    if path.name.startswith("~$"):
        raise ValueError("Excel 임시 파일 대신 원래 파일을 선택해 주세요.")
    if not path.is_file():
        raise ValueError("파일을 찾을 수 없습니다. 파일 위치를 확인해 주세요.")
    return path.resolve()


def select_dropped_file(tk, data):
    # Explorer supplies a Tcl list, including braces for spaces and Korean paths.
    paths = tk.splitlist(data)
    if len(paths) != 1:
        raise ValueError("Excel 파일을 한 번에 하나씩 놓아 주세요.")
    return validate_excel_path(paths[0])


def _excel_rows(path):
    if path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book.worksheets[0]
            # Some exporters report incorrect bounds; match the previous Excel reader.
            sheet.reset_dimensions()
            for row in sheet.iter_rows():
                yield [None if cell.data_type == "e" else cell.value for cell in row]
        finally:
            book.close()
    else:
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            sheet = book.sheet_by_index(0)
            for row in range(sheet.nrows):
                values = []
                for cell in sheet.row(row):
                    value = cell.value
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            parsed = xlrd.xldate_as_datetime(value, book.datemode)
                            anchor = (1904, 1, 1) if book.datemode else (1899, 12, 31)
                            value = parsed.time() if (parsed.year, parsed.month, parsed.day) == anchor else parsed
                        except OverflowError:
                            pass
                    values.append(value)
                yield values
        finally:
            book.release_resources()


def _label_text(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    # A cell is one label line. Control characters must not enter DXF tags.
    return " ".join(text.splitlines())


def extract_blocks_auto(excel_path):
    columns, active, boolean_numbers = [], [], []
    for row in _excel_rows(validate_excel_path(excel_path)):
        while len(columns) < len(row):
            columns.append([])
            active.append([])
            boolean_numbers.append({})
        for column in range(len(columns)):
            value = row[column] if column < len(row) else None
            # The former object reader interns equal bool/0/1 values per column.
            # Preserve its displayed labels even when those Excel types are mixed.
            if isinstance(value, (bool, int, float)) and value in (0, 1):
                value = boolean_numbers[column].setdefault(value, value)
            text = _label_text(value)
            if text:
                active[column].append(text)
            elif active[column]:
                columns[column].append(active[column])
                active[column] = []
    for column, block in enumerate(active):
        if block:
            columns[column].append(block)
    return [block for column in columns for block in column]


@dataclass(frozen=True)
class LabelSettings:
    cell_w: float = 60
    cell_h: float = 15
    gap_x: float = 20
    gap_y: float = 20
    cols_per_row: int = 5
    text_scale: float = 0.45
    auto_fit: bool = True
    font_key: str = DEFAULT_FONT
    box_color: int = 7
    text_color: int = 7
    min_text_height: float = 2.5

    def __post_init__(self):
        for name in ("cell_w", "cell_h", "text_scale", "gap_x", "gap_y", "min_text_height"):
            value = getattr(self, name)
            if type(value) not in (int, float) or not math.isfinite(value):
                raise ValueError("규격에는 유한한 숫자를 입력해 주세요.")
            if value < 0 or (name in {"cell_w", "cell_h", "text_scale", "min_text_height"} and value == 0):
                raise ValueError("폭, 높이와 글자 비율은 0보다 커야 하며 간격은 0 이상이어야 합니다.")
        if type(self.cols_per_row) is not int or self.cols_per_row < 1:
            raise ValueError("한 행의 라벨 수는 1 이상의 정수여야 합니다.")
        if self.font_key not in CAD_FONTS:
            raise ValueError("목록에서 글꼴을 선택해 주세요.")
        if any(type(c) is not int or c not in range(1, 8)
               for c in (self.box_color, self.text_color)):
            raise ValueError("목록에서 도면 색상을 선택해 주세요.")
        if not all(math.isfinite(v) for v in (self.cell_w + self.gap_x, self.cell_h + self.gap_y, self.cell_h * self.text_scale)):
            raise ValueError("규격의 합과 글자 높이가 너무 큽니다. 폭, 높이와 간격을 줄여 주세요.")
        if self.cell_h * self.text_scale <= 0:
            raise ValueError("계산된 글자 높이가 너무 작습니다. 높이와 글자 비율을 늘려 주세요.")


def get_text_display_units(text):
    return max(sum(1.15 if unicodedata.east_asian_width(ch) in {"F", "W"}
                   else 0.70 if ch.isupper() or ch.isdigit() else 0.55
                   for ch in text), 1.0)


@dataclass(frozen=True)
class PlacedCell:
    label: int
    text: str
    corners: tuple
    center: tuple
    text_height: float
    width_factor: float


@dataclass(frozen=True)
class Layout:
    settings: LabelSettings
    cells: tuple
    bounds: tuple
    label_bounds: tuple
    warnings: tuple


def compute_layout(blocks, settings, token=None):
    row_heights = {}
    for index, block in enumerate(blocks):
        if token:
            token.check()
        row = index // settings.cols_per_row
        row_heights[row] = max(row_heights.get(row, 0), len(block) * settings.cell_h)
    offsets, offset = {}, 0
    for row, height in row_heights.items():
        offsets[row] = offset
        offset += height + settings.gap_y
        if not math.isfinite(offset):
            raise ValueError("전체 배치 높이가 너무 큽니다. 칸 높이와 세로 간격을 줄여 주세요.")
    cells, bounds, warnings = [], [], []
    for index, block in enumerate(blocks):
        if token:
            token.check()
        x = index % settings.cols_per_row * (settings.cell_w + settings.gap_x)
        y = -offsets[index // settings.cols_per_row]
        if not all(math.isfinite(v) for v in (x, y, x + settings.cell_w, y - len(block) * settings.cell_h)):
            raise ValueError("전체 배치 좌표가 너무 큽니다. 폭, 높이와 간격 또는 배치 개수를 줄여 주세요.")
        bounds.append((x, y - len(block) * settings.cell_h, x + settings.cell_w, y))
        for line, text in enumerate(block):
            top = y - line * settings.cell_h
            bottom = top - settings.cell_h
            height, width = settings.cell_h * settings.text_scale, 1.0
            if settings.auto_fit:
                denominator = get_text_display_units(text) * height * 1.2
                if not math.isfinite(denominator) or denominator <= 0:
                    raise ValueError("문자 크기를 계산할 수 없습니다. 글자 높이와 문구 길이를 확인해 주세요.")
                ratio = settings.cell_w * 0.75 / denominator
                if ratio < 1:
                    width = max(ratio, 0.65)
                    if ratio < 0.65:
                        height *= ratio / 0.65
                height = min(height, settings.cell_h * 0.75)
            if not math.isfinite(height) or height <= 0 or not math.isfinite(width):
                raise ValueError("계산된 글자 높이나 배율이 유효하지 않습니다. 규격을 확인해 주세요.")
            if height < settings.min_text_height:
                warnings.append((index, "small-text", f"글자 높이 {height:.3f} mm가 최소 {settings.min_text_height:g} mm보다 작습니다. 폭이나 문구 및 칸 분리를 확인해 주세요."))
            if get_text_display_units(text) * height * 1.2 * width > settings.cell_w * 0.95 or height > settings.cell_h * 0.9:
                warnings.append((index, "overflow", "글자가 칸을 넘을 수 있습니다. 폭이나 높이 및 자동 맞춤을 확인해 주세요."))
            center = (x + settings.cell_w / 2, top + (bottom - top) / 2)
            if not all(math.isfinite(v) for v in (*center, top, bottom, height, width)):
                raise ValueError("문자의 좌표와 크기를 계산할 수 없습니다. 규격을 줄여 주세요.")
            cells.append(PlacedCell(index, text, ((x, top), (x + settings.cell_w, top), (x + settings.cell_w, bottom), (x, bottom)), center, height, width))
    overall = (0, min((b[1] for b in bounds), default=0), max((b[2] for b in bounds), default=0), 0)
    return Layout(settings, tuple(cells), overall, tuple(bounds), tuple(warnings))


def drawing_from_layout(layout, token=None):
    settings = layout.settings
    doc = ezdxf.new("R2010")  # AC1024, Unicode DXF written as UTF-8.
    doc.units = ezdxf.units.MM
    style, font = CAD_FONTS[settings.font_key]
    if style == "STANDARD":
        doc.styles.get(style).dxf.font = font
    else:
        doc.styles.new(style, dxfattribs={"font": font})
    doc.layers.new("LABEL_BOX", dxfattribs={"color": settings.box_color})
    doc.layers.new("TEXT_DATA", dxfattribs={"color": settings.text_color})
    model = doc.modelspace()
    for cell in layout.cells:
        if token:
            token.check()
        model.add_lwpolyline(cell.corners, close=True, dxfattribs={"layer": "LABEL_BOX"})
        model.add_text(cell.text, dxfattribs={"layer": "TEXT_DATA", "style": style,
                                          "height": cell.text_height, "width": cell.width_factor}).set_placement(
            cell.center, align=TextEntityAlignment.MIDDLE_CENTER)
    return doc


def create_drawing(blocks, settings):
    return drawing_from_layout(compute_layout(blocks, settings))


def generate_dxf_string(excel_path, cell_w=60, cell_h=15, gap_x=20, gap_y=20,
                        cols_per_row=5, text_scale=0.45, auto_fit=True,
                        font_key=DEFAULT_FONT, box_color=7, text_color=7):
    settings = LabelSettings(cell_w, cell_h, gap_x, gap_y, cols_per_row,
                             text_scale, auto_fit, font_key, box_color, text_color)
    blocks = extract_blocks_auto(excel_path)
    if not blocks:
        return "", 0
    stream = StringIO()
    create_drawing(blocks, settings).write(stream)
    return stream.getvalue(), len(blocks)


def save_drawing(doc, output, token=None):
    output = Path(output)
    if output.suffix.lower() != ".dxf":
        raise ValueError("저장 파일의 확장자는 .dxf여야 합니다.")
    temp = None
    try:
        if token:
            token.check()
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="\n",
                                         dir=output.parent, suffix=".tmp", delete=False) as stream:
            temp = Path(stream.name)
            doc.write(stream)
        if token:
            token.commit(lambda: os.replace(temp, output))
        else:
            os.replace(temp, output)
    finally:
        if temp is not None and temp.exists():
            temp.unlink()
