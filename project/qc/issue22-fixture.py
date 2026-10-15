"""Create synthetic, text-formatted Excel roundtrip fixtures. No user files."""
import argparse
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import Workbook

ROWS = [
    ["AAA", "BBB", "CCC"],
    ["DDD", "EEE", ""],
    ["FFF", "", "GGG"],
    ["", "선두 빈 칸", "한글"],
    ["001", "12345678901234567890", "2026-01-02"],
    ["=SUM(A1:A2)", '따옴표 "확인"', " 공백 보존 "],
]


def create(directory):
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / "issue22-roundtrip.xlsx"
    if path.exists():
        raise FileExistsError("Use a new evidence directory; keep existing fixtures.")
    book = Workbook()
    source = book.active
    source.title = "Source"
    destination = book.create_sheet("Roundtrip")
    for sheet in (source, destination):
        for row in sheet.iter_rows(min_row=1, max_row=20, min_col=1, max_col=6):
            for cell in row:
                cell.number_format = "@"
        for column in ("A", "B", "C", "D", "E", "F"):
            sheet.column_dimensions[column].width = 26
    for r, row in enumerate(ROWS, 1):
        for c, value in enumerate(row, 1):
            cell = source.cell(r, c, value)
            cell.data_type = "s"
    book.save(path)
    book.close()
    manifest = {"fixture": path.name, "sha256": sha256(path.read_bytes()).hexdigest(),
                "source_sheet": "Source", "destination_sheet": "Roundtrip",
                "range": "A1:C6", "rows": ROWS, "destination_number_format": "@"}
    (directory / "fixture.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"fixture": str(path), "sha256": manifest["sha256"]}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--directory", required=True, type=Path)
    create(parser.parse_args().directory)
